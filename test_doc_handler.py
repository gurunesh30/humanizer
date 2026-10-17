import sys
import tempfile
from pathlib import Path
from docx import Document

# Ensure workspace root is in python path
sys.path.insert(0, str(Path(__file__).parent))

from doc_handler import DocumentHandler


def test_document_handler_read_and_filter():
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = Path(temp_dir) / "test_input.docx"
        doc = Document()
        doc.add_paragraph("First paragraph content.")
        doc.add_paragraph("   ")  # Whitespace only
        doc.add_paragraph("")      # Blank
        doc.add_paragraph("Second non-empty paragraph.")
        doc.save(str(test_file))

        handler = DocumentHandler(input_path=str(test_file))
        paragraphs = handler.read_paragraphs()

        assert len(paragraphs) == 2
        assert paragraphs[0] == "First paragraph content."
        assert paragraphs[1] == "Second non-empty paragraph."


def test_document_handler_incremental_append():
    with tempfile.TemporaryDirectory() as temp_dir:
        out_file = Path(temp_dir) / "test_output.docx"
        handler = DocumentHandler(output_path=str(out_file))

        handler.init_output_document()
        assert out_file.exists()

        handler.append_paragraph("Paraphrased 1")
        # Read back immediately to verify persistence
        d1 = Document(str(out_file))
        assert len(d1.paragraphs) == 1
        assert d1.paragraphs[0].text == "Paraphrased 1"

        handler.append_paragraph("Paraphrased 2")
        d2 = Document(str(out_file))
        assert len(d2.paragraphs) == 2
        assert d2.paragraphs[1].text == "Paraphrased 2"


if __name__ == "__main__":
    test_document_handler_read_and_filter()
    test_document_handler_incremental_append()
    print("All doc_handler tests passed!")
