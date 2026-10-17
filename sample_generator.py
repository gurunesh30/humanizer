"""Generate sample input documents for testing the paraphrasing automation pipeline."""

from pathlib import Path
from docx import Document


def create_sample_document(output_path: str = "sample_input.docx") -> Path:
    """Create a sample Word document with test paragraphs and spacing.

    Args:
        output_path: Target filename for sample document.

    Returns:
        Path to the created document.
    """
    doc = Document()
    doc.add_heading("ZeroGPT Automation Test Document", level=1)

    doc.add_paragraph(
        "Artificial intelligence and automated language modeling systems have drastically revolutionized "
        "the way people conduct academic research and author creative content across diverse industries."
    )

    # Empty paragraph to test whitespace filtering
    doc.add_paragraph("   ")

    doc.add_paragraph(
        "When constructing automation pipelines for web-based services, developers must consider rate limiting, "
        "transient network anomalies, and anti-scraping defenses like CAPTCHAs and Cloudflare Turnstile."
    )

    doc.add_paragraph(
        "Incremental document persistence guarantees that processed data is immediately written to non-volatile "
        "storage, mitigating the risk of total data loss if an unexpected interruption occurs."
    )

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    print(f"Sample test document created at: {target.resolve()}")
    return target


if __name__ == "__main__":
    create_sample_document()
